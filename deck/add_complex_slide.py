"""Insert the "Complex Software" problem slide as slide 3.

Replaces the four-panel, seventeen-question layout with one graphic that
carries the argument: a progress bar where the part an AI tool delivers is a
sliver and the four engineering categories fill the rest, followed by the
cost chain that the original slide buried at the bottom edge.

It stays a problem slide — no resolution line, no bridge to the solution.

Runs against the built deck (the original export is no longer available), so
it also renumbers the section markers of every slide it pushes down.
"""
import zipfile

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DECK = "Pitch_Deck_AthenaRun_v2.pptx"
W, H, PX = 1920, 1080, 6350

BG     = RGBColor(0x0E, 0x0E, 0x11)
WHITE  = RGBColor(0xEC, 0xEC, 0xEE)
MUTED  = RGBColor(0x9A, 0x9A, 0xA2)
ORANGE = RGBColor(0xFF, 0x6B, 0x3D)
RULE   = RGBColor(0x2B, 0x2B, 0x2D)
PANEL  = RGBColor(0x16, 0x16, 0x19)
FILLED = RGBColor(0x3A, 0x3A, 0x3E)

SANS, MONO = "Poppins", "Courier New"
TITLE, SUB, MARKER, SECLABEL = 27, 14, 11, 11
L, R = 110, 1809


def shape(sl, kind, x, y, w, h, fill, line_col=None, line_w=1.25, rot=None):
    s = sl.shapes.add_shape(kind, int(x) * PX, int(y) * PX,
                            max(int(w), 1) * PX, max(int(h), 1) * PX)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line_col is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_col; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if rot is not None:
        s.rotation = rot
    return s


def rect(sl, x, y, w, h, colour, line_col=None):
    return shape(sl, MSO_SHAPE.RECTANGLE, x, y, w, h, colour, line_col)


def line(sl, x, y, w, h, parts, align=PP_ALIGN.LEFT, spacing=None):
    tb = sl.shapes.add_textbox(int(x) * PX, int(y) * PX, int(w) * PX, int(h) * PX)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    if spacing:
        p.line_spacing = Pt(spacing)
    for text, size, col, font, bold, track in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size); r.font.color.rgb = col
        r.font.name = font; r.font.bold = bold
        if track:
            r.font._rPr.set("spc", str(int(track * 100)))
    return tf


def arrowhead(sl, x, y, colour, size=14):
    shape(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, x - size / 2, y - size / 2,
          size, size, colour, rot=90)


prs = Presentation(DECK)
slides = list(prs.slides)

# ---- renumber the markers of everything that moves down ------------------
for sl in slides[2:10]:
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                t = r.text
                if len(t) < 32 and " — " in t and t[:2].isdigit():
                    r.text = f"{int(t[:2]) + 1:02d}{t[2:]}"

# ---- the new slide -------------------------------------------------------
sc = prs.slides.add_slide(slides[1].slide_layout)
sc.shapes.add_picture("bg.png", 0, 0, Emu(W * PX), Emu(H * PX))

line(sc, L, 90, 1560, 62,
     [("Complex Software.", TITLE, WHITE, SANS, True, None),
      (" Is it possible with AI tools?", TITLE, MUTED, SANS, False, None)])
line(sc, R - 500, 98, 500, 24,
     [("03 — The illusion", MARKER, MUTED, MONO, False, 1.2)],
     align=PP_ALIGN.RIGHT)
rect(sc, L, 172, R - L, 1, RULE)

line(sc, L, 196, 1520, 76,
     [("Starting has never been cheaper: an AI IDE, a chat, a prompt, and "
       "something runs. That ease is what makes the engineering behind it "
       "easy to underestimate.", SUB, MUTED, SANS, False, None)], spacing=19)

# ---- the bar: what a tool delivers, against what engineering still owes ---
SEG1 = 204
SEGW = (R - L - SEG1) / 4
edges = [L, L + SEG1] + [L + SEG1 + i * SEGW for i in range(1, 5)]

line(sc, L, 322, 400, 26,
     [("DAY ONE", SECLABEL, MUTED, MONO, False, 2.5)])
line(sc, edges[1], 322, 900, 26,
     [("EVERYTHING AFTER", SECLABEL, MUTED, MONO, False, 2.5)])

rect(sc, L, 370, SEG1, 80, FILLED)
line(sc, L, 370, SEG1, 80,
     [("Prototype", 16, WHITE, SANS, False, None)], align=PP_ALIGN.CENTER)
line(sc, L, 466, SEG1, 34,
     [("in an afternoon", 12.5, MUTED, SANS, False, None)], align=PP_ALIGN.CENTER)

CATS = [("Fit", "Solves the real process?"),
        ("Integration", "Lives in your estate?"),
        ("Assurance", "Someone can sign it?"),
        ("Operation", "Survives year two?")]
# one continuous block, subdivided — so the eye reads "sliver vs. the rest"
rect(sc, edges[1], 370, R - edges[1], 80, PANEL, RULE)
for i in range(1, 4):
    rect(sc, edges[1 + i], 378, 1, 64, RULE)
for i, (name, question) in enumerate(CATS):
    x = edges[1 + i]
    line(sc, x, 370, SEGW, 80,
         [(name, 16, WHITE, SANS, False, None)], align=PP_ALIGN.CENTER)
    line(sc, x, 466, SEGW, 34,
         [(question, 12.5, MUTED, SANS, False, None)], align=PP_ALIGN.CENTER)

rect(sc, L, 566, R - L, 1, RULE)

# ---- the cost chain, no longer a footnote --------------------------------
line(sc, L, 598, 900, 26,
     [("WHAT THE GAP COSTS", SECLABEL, MUTED, MONO, False, 2.5)])
STEPS = ["Easy start", "Underestimated engineering", "Hidden work", "Rework",
         "Lost speed, higher cost", "Lost EBIT"]
pitch = (R - L) / 6
for i, step in enumerate(STEPS):
    cx = L + pitch / 2 + i * pitch
    last = i == len(STEPS) - 1
    line(sc, cx - 125, 646, 250, 76,
         [(step, 14, ORANGE if last else WHITE, SANS, last, None)],
         align=PP_ALIGN.CENTER, spacing=19)
    if i:
        arrowhead(sc, L + i * pitch, 684, MUTED)

rect(sc, L, 790, R - L, 1, RULE)
line(sc, L, 822, 1600, 38,
     [("None of this is optional.", 16, WHITE, SANS, True, None),
      ("  It only gets discovered late.", 16, MUTED, SANS, False, None)])

sc.shapes.add_picture("logo.png", 100 * PX, 1014 * PX, 170 * PX, 40 * PX)

# ---- move it into position 3 --------------------------------------------
lst = prs.slides._sldIdLst
ids = list(lst)
current = list(prs.slides)
order = current[:2] + [sc] + current[2:-1]
for e in ids:
    lst.remove(e)
for sl in order:
    lst.append(ids[current.index(sl)])

prs.save(DECK)
print(f"wrote {DECK}: {len(prs.slides)} slides")
