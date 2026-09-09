"""Traction & roadmap — native slide, second to last.

Left: the pipeline as three stacked figures plus the partners and programmes
backing the company. Right: the 90-day plan and what follows. Facts marked
CONFIRM in the reply come from the previous traction/milestones slides and
must be re-validated by the founders before this goes out.

Outputs: an import file in Helvetica for Keynote, and the slide replacing the
two old image slides (traction, milestones) at the end of the PPTX deck.
"""
import re

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from native_slides import (F, W, H, PX, L, R, WHITE, MUTED, DIM, ORANGE, RULE,
                           DIV, PANEL, S, M, text, rect, header, logo, background,
                           _new_deck, _drop_slide)


def traction(sl):
    background(sl)
    header(sl, "Traction & roadmap.", "Demand before launch.", "TRACTION")
    text(sl, L, 196, 1560, 34,
         [S("Signed intent, a bank contract in signing and a Q4 closing pipeline — "
            "then ninety days to convert it.", 14, MUTED)])

    # ---- left: pipeline -------------------------------------------------
    text(sl, L, 292, 600, 24, [M("PIPELINE", 10.5)])
    rows = [("~€50K", ORANGE, "International bank", "contract in signing"),
            ("€10–20K", WHITE, "per deal · several deals", "closing in Q4 2026"),
            ("€150K", WHITE, "signed LOI pipeline", "7 letters of intent")]
    for i, (num, col, l1, l2) in enumerate(rows):
        y = 330 + i * 84
        text(sl, L, y, 300, 60, [S(num, 34, col)])
        text(sl, 430, y + 4, 560, 28, [S(l1, 15, WHITE)])
        text(sl, 430, y + 32, 560, 24, [S(l2, 12, MUTED)])
        if i < 2:
            rect(sl, L, y + 72, 880, 1, DIV)
    rect(sl, L, 590, 880, 1, RULE)

    text(sl, L, 614, 600, 24, [M("BACKED & DISTRIBUTED", 10.5)])
    backers = [(L,   648, "Channel partner",       "DISTRIBUTION"),
               (560, 648, "Hessen AI Cohort 2026", "STATE ACCELERATOR"),
               (L,   722, "Cyberlab Cohort 2026",  "ACCELERATOR · KARLSRUHE"),
               (560, 722, "NVIDIA Inception",      "INFRASTRUCTURE")]
    for x, y, name, sub in backers:
        text(sl, x, y, 430, 28, [S(name, 14, WHITE)])
        text(sl, x, y + 30, 430, 22, [M(sub, 9.5, MUTED, 1.8)])

    # ---- right: roadmap -------------------------------------------------
    rect(sl, 1040, 286, 1, 520, DIV)
    X, WD = 1090, R - 1090
    phases = [
        (292, "NEXT 90 DAYS  ·  Q4 2026", "Prove & launch",
         ["Close the bank contract and the Q4 deals",
          "Convert the signed LOIs into paying pilots",
          "Platform launch — go-to-market live"]),
        (556, "Q1 – Q2 2027", "Integrate & scale",
         ["Ship integrations across CRM, ERP and Slack",
          "Scale through the channel partner and accelerators",
          "Open the app marketplace, enterprise partnerships"]),
    ]
    for y0, label, head, bullets in phases:
        text(sl, X, y0, WD, 24, [M(label, 10.5, ORANGE if y0 == 292 else MUTED)])
        text(sl, X, y0 + 30, WD, 36, [S(head, 20, WHITE)])
        for i, b in enumerate(bullets):
            y = y0 + 80 + i * 50
            text(sl, X, y, 20, 30, [S("–", 12.5, DIM)], anchor=MSO_ANCHOR.TOP)
            text(sl, X + 26, y, WD - 26, 46, [S(b, 12.5, MUTED)], spacing=16,
                 anchor=MSO_ANCHOR.TOP)
    rect(sl, X, 534, WD, 1, RULE)

    rect(sl, L, 840, R - L, 1, RULE)
    text(sl, L, 856, 1680, 40,
         [S("The hard part is done.", 16, WHITE, True),
          S(" Contracts in signing, engine shipped — what’s ahead is distribution, "
            "not invention.", 16, MUTED)])
    logo(sl)


def build_import(path):
    F["sans"] = "Helvetica"
    prs, blank = _new_deck()
    traction(prs.slides.add_slide(blank))
    prs.save(path)
    print(f"wrote {path} (Helvetica)")


def update_deck(path):
    F["sans"] = "Poppins"
    prs = Presentation(path)
    slides = list(prs.slides)
    old = [s for s in slides if all(sh.shape_type == 13 for sh in s.shapes)]
    assert len(old) == 2, f"expected the two remaining image slides, found {len(old)}"
    new = prs.slides.add_slide(slides[1].slide_layout)
    traction(new)
    lst = prs.slides._sldIdLst
    ids = list(lst); cur = list(prs.slides)
    order = [s for s in cur if s not in old]          # new slide is already last
    for e in ids: lst.remove(e)
    for s in order: lst.append(ids[cur.index(s)])
    # drop the old slides — and any orphaned slide rels left by earlier edits,
    # which otherwise collide with the renumbered part names on save
    live = {sid.rId for sid in lst}
    for rId, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith("/slide") and rId not in live:
            prs.part.rels.pop(rId)
    prs.save(path)
    print(f"wrote {path}: {len(prs.slides)} slides (Poppins)")


if __name__ == "__main__":
    build_import("AthenaRun_slide_traction_native.pptx")
    update_deck("Pitch_Deck_AthenaRun_v2.pptx")
