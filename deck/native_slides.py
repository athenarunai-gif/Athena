"""Native rebuilds of the three remaining image slides: market, landscape, team.

Every element is its own shape or text box so it can be selected and edited.
Coordinates are in the deck's 1920x1080 export grid (1 px = 6350 EMU).

Two outputs:
  * AthenaRun_slides_10-12_native.pptx  -- Helvetica, for import into Keynote
  * Pitch_Deck_AthenaRun_v2.pptx        -- Poppins, the three image slides
                                            replaced in place; markers unified
                                            (mono caps, no numbers); lifecycle
                                            slide says multiple agents
"""
import copy
import re
import sys

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

W, H, PX = 1920, 1080, 6350
BG     = RGBColor(0x0E, 0x0E, 0x11)
WHITE  = RGBColor(0xEC, 0xEC, 0xEE)
MUTED  = RGBColor(0x9A, 0x9A, 0xA2)
DIM    = RGBColor(0x5A, 0x5A, 0x60)
ORANGE = RGBColor(0xFF, 0x6B, 0x3D)
RULE   = RGBColor(0x2B, 0x2B, 0x2D)
DIV    = RGBColor(0x1F, 0x1F, 0x22)
PANEL  = RGBColor(0x16, 0x16, 0x19)
MONO = "Courier New"
L, R = 110, 1809

F = {"sans": "Poppins"}          # set by the driver per output


def shape(sl, kind, x, y, w, h, fill, line_col=None, line_w=1.25):
    s = sl.shapes.add_shape(kind, int(x) * PX, int(y) * PX,
                            max(int(w), 1) * PX, max(int(h), 1) * PX)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line_col is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_col; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s


def rect(sl, x, y, w, h, colour, line_col=None):
    return shape(sl, MSO_SHAPE.RECTANGLE, x, y, w, h, colour, line_col)


def text(sl, x, y, w, h, parts, align=PP_ALIGN.LEFT, spacing=None,
         anchor=MSO_ANCHOR.MIDDLE):
    tb = sl.shapes.add_textbox(int(x) * PX, int(y) * PX, int(w) * PX, int(h) * PX)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if spacing:
        p.line_spacing = Pt(spacing)
    for t, size, col, mono, bold, track in parts:
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size); r.font.color.rgb = col
        r.font.name = MONO if mono else F["sans"]; r.font.bold = bold
        if track:
            r.font._rPr.set("spc", str(int(track * 100)))
    return tf


def S(t, size, col=WHITE, bold=False):      # sans run
    return (t, size, col, False, bold, None)


def M(t, size=11, col=MUTED, track=2.5):    # mono label run
    return (t, size, col, True, False, track)


def header(sl, bold_part, muted_part, marker, size=27):
    text(sl, L, 90, 1480, 62, [S(bold_part, size, WHITE, True), S(" " + muted_part, size, MUTED)])
    text(sl, R - 500, 98, 500, 24, [M(marker, 11, MUTED, 1.2)], align=PP_ALIGN.RIGHT)
    rect(sl, L, 172, R - L, 1, RULE)


def logo(sl):
    sl.shapes.add_picture("logo.png", 100 * PX, 1014 * PX, 170 * PX, 40 * PX)


def background(sl):
    sl.shapes.add_picture("bg.png", 0, 0, Emu(W * PX), Emu(H * PX))


# ---------------------------------------------------------------------------
def market(sl):
    background(sl)
    header(sl, "The market.", "Three cuts, each one narrower.", "MARKET SIZE")
    text(sl, L, 196, 1400, 34, [S("Every number below carries its assumption.", 14, MUTED)])
    rect(sl, L, 384, R - L, 1, RULE)
    cols = [
        (L,    "TAM", WHITE,  "CUSTOM SOFTWARE DELIVERY, GLOBAL", "~$65B",   "2026",
         "Analyst range $51–74B. 21% CAGR through 2030."),
        (735,  "SAM", WHITE,  "EUROPE  ·  REGULATED SECTORS",      "$15–20B", "serviceable",
         "Delivery requiring a compliance gate and an audit trail."),
        (1300, "SOM", ORANGE, "5 YEARS  ·  CUMULATIVE CONTRACT VALUE", "$150M", "TCV",
         "250 customers × $600k. Both assumptions."),
    ]
    for x, name, col, sub, big, suffix, note in cols:
        text(sl, x, 424, 500, 34, [S(name, 17, col, True)])
        text(sl, x, 462, 520, 24, [M(sub, 10.5)])
        text(sl, x, 500, 500, 96, [S(big, 40, WHITE), S(" " + suffix, 12, MUTED)])
        rect(sl, x, 608, 500, 1, RULE)
        text(sl, x, 622, 460, 72, [S(note, 13, MUTED)], spacing=18, anchor=MSO_ANCHOR.TOP)
    for x in (677, 1243):
        rect(sl, x, 420, 1, 262, DIV)
    rect(sl, L, 736, R - L, 100, PANEL)
    text(sl, 155, 736, 110, 100, [M("AMBITION", 10.5)])
    text(sl, 275, 736, 1400, 100, [S("1% of SAM", 17, WHITE, True),
                                   S(" — open source as the adoption lever.", 17, MUTED)])
    logo(sl)


# ---------------------------------------------------------------------------
def _mark(sl, cx, cy, kind, col):
    """full = filled circle, partial = hollow circle, none = short dash."""
    if kind == "full":
        shape(sl, MSO_SHAPE.OVAL, cx - 7, cy - 7, 14, 14, col)
    elif kind == "partial":
        shape(sl, MSO_SHAPE.OVAL, cx - 7, cy - 7, 14, 14, None, col, 1.5)
    else:
        rect(sl, cx - 6, cy - 1, 12, 2, DIM)


def landscape(sl):
    background(sl)
    header(sl, "The landscape.", "Everyone owns a segment. We own the chain.", "COMPETITIVE MAP", size=23)
    text(sl, L, 190, 1100, 76,
         [S("Coding agents write. App builders assemble. Regulated-agent platforms "
            "automate workflows. None run idea to running software end to end.", 14, MUTED)],
         spacing=19)
    rect(sl, 716, 292, 302, 524, PANEL)                    # AthenaRun column band
    heads = [(740, "AthenaRun", ORANGE, "idea to operate"),
             (1018, "Coding agents", WHITE, "Devin, Copilot"),
             (1283, "App builders", WHITE, "Reflex, Bubble"),
             (1546, "Reg. agent platforms", WHITE, "Palantir, MightyBot")]
    for x, name, col, sub in heads:
        text(sl, x, 300, 262, 78, [S(name, 15, col)], spacing=19, anchor=MSO_ANCHOR.BOTTOM)
        text(sl, x, 382, 262, 24, [M(sub, 10.5, MUTED, 1.2)])
    rows = [("Requirements from conversation", "full", "none",    "none",    "partial"),
            ("Compliance advisory before code", "full", "none",    "partial", "full"),
            ("Autonomous build + review",       "full", "full",    "full",    "partial"),
            ("Human gate at every release",     "full", "partial", "partial", "full"),
            ("Operate after go-live",           "full", "none",    "partial", "partial"),
            ("Learning loop across projects",   "full", "none",    "none",    "none")]
    xs = [867, 1151, 1414, 1678]
    for i, (label, *marks) in enumerate(rows):
        y = 410 + i * 68
        rect(sl, L, y, R - L, 1, RULE)
        text(sl, L, y + 1, 580, 67, [S(label, 15, WHITE)])
        for j, (cx, kind) in enumerate(zip(xs, marks)):
            _mark(sl, cx, y + 34, kind, ORANGE if j == 0 else MUTED)
    rect(sl, L, 818, R - L, 1, RULE)
    _mark(sl, 117, 847, "full", MUTED);    text(sl, 130, 835, 80, 24, [M("full", 10.5, MUTED, 1.0)])
    _mark(sl, 205, 847, "partial", MUTED); text(sl, 218, 835, 100, 24, [M("partial", 10.5, MUTED, 1.0)])
    _mark(sl, 322, 847, "none", MUTED);    text(sl, 334, 835, 80, 24, [M("none", 10.5, MUTED, 1.0)])
    rect(sl, L, 886, R - L, 1, RULE)
    text(sl, L, 896, 1680, 80, [S("The others compete on one column.", 16, MUTED),
                                S(" Our moat is the whole row — and the loop nobody else has.", 16, WHITE, True)],
         spacing=22)
    logo(sl)


# ---------------------------------------------------------------------------
def team(sl):
    background(sl)
    header(sl, "Built by people who’ve", "done both sides.", "TEAM")
    text(sl, L, 190, 1000, 76,
         [S("One sold enterprise software. One built it. Together they’ve seen "
            "the pain from both sides.", 14, MUTED)], spacing=19)
    people = [
        (L, 770, "GO-TO-MARKET", "Patrick Knapp", "CEO",
         ["10+ years in B2B sales & GTM across Seed–Series E startups",
          "Closed mid-market and enterprise deals at 6–7 figure ACVs",
          "Record €7M enterprise landing deal closed",
          "Scaled business functions to 20+"]),
        (1040, 769, "PRODUCT & ENGINEERING", "Michael Altermann", "CTO",
         ["10+ years in product development and software engineering",
          "Built and scaled a core logistics platform for daily operations",
          "Scaled engineering teams to 20+",
          "Owns architecture and the autonomous build pipeline"]),
    ]
    for x, w, area, name, role, bullets in people:
        text(sl, x, 376, w, 24, [M(area, 10.5)])
        rect(sl, x, 414, w, 1, RULE)
        text(sl, x, 444, w, 48, [S(name, 24, WHITE), S("   ", 24), M(role, 10.5, MUTED, 1.5)])
        rect(sl, x, 520, w, 1, RULE)
        for i, b in enumerate(bullets):
            y = 540 + i * 70
            text(sl, x, y, 20, 30, [S("–", 12, DIM)], anchor=MSO_ANCHOR.TOP)
            text(sl, x + 28, y, w - 28, 58, [S(b, 12, MUTED)], spacing=15, anchor=MSO_ANCHOR.TOP)
            if i < 3:
                rect(sl, x, y + 60, w, 1, DIV)
    rect(sl, 966, 380, 1, 420, DIV)
    rect(sl, L, 826, R - L, 1, RULE)
    text(sl, L, 838, 1680, 80, [S("+ Christian Lux, Co-Founder, Full-Stack Developer", 16, WHITE, True),
                                S(" — highly experienced in AI development and security.", 16, MUTED)],
         spacing=22)
    logo(sl)


# ---------------------------------------------------------------------------
def _new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    return prs, prs.slide_layouts[6]


def build_import_file(path):
    F["sans"] = "Helvetica"
    prs, blank = _new_deck()
    for fn in (market, landscape, team):
        fn(prs.slides.add_slide(blank))
    prs.save(path)
    print(f"wrote {path}: {len(prs.slides)} slides (Helvetica)")


def _drop_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.slides.part.related_part(sid.rId) is slide.part if hasattr(prs.slides.part, "related_part") else False:
            pass
    for sid in list(lst):
        if sid.rId in prs.part.rels and prs.part.rels[sid.rId].target_part is slide.part:
            prs.part.drop_rel(sid.rId); lst.remove(sid); return


def update_main_deck(path):
    F["sans"] = "Poppins"
    prs = Presentation(path)
    slides = list(prs.slides)
    blank = slides[1].slide_layout
    # market / landscape / team are the three image-only slides after the divider
    old = [s for s in slides if all(sh.shape_type == 13 for sh in s.shapes)][0:3]
    assert len(old) == 3, "expected the three image appendix slides"
    new = []
    for fn in (market, landscape, team):
        s = prs.slides.add_slide(blank); fn(s); new.append(s)
    lst = prs.slides._sldIdLst
    ids = list(lst); cur = list(prs.slides)
    order = []
    for s in cur:
        if s in old: order.append(new[old.index(s)])
        elif s in new: continue
        else: order.append(s)
    for e in ids: lst.remove(e)
    for s in order: lst.append(ids[cur.index(s)])
    for s in old: _drop_slide(prs, s)
    # markers: mono caps, no numbers; lifecycle: multiple agents
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    m = re.fullmatch(r"(\d\d) — (.+)", r.text.strip())
                    if m: r.text = m.group(2).upper()
                    if r.text == "ONE AGENT PER PHASE": r.text = "MULTIPLE AGENTS PER PHASE"
                    if r.text == "Every phase has an agent.": r.text = "Every phase has multiple agents."
                    if r.text == " Every phase has an agent.": r.text = " Every phase has multiple agents."
                    if re.fullmatch(r"(Planning|Analysis|Design|Build|Test|Operations) agent", r.text):
                        r.text += "s"
    prs.save(path)
    print(f"wrote {path}: {len(prs.slides)} slides (Poppins)")


if __name__ == "__main__":
    build_import_file("AthenaRun_slides_10-12_native.pptx")
    update_main_deck("Pitch_Deck_AthenaRun_v2.pptx")
