"""Replace the three stat blocks and the source footer on slide 2.

Slide 2 is a full-bleed PNG export. The old stats and the old source line are
erased from the image itself (so the superseded citations are really gone, not
just covered), and the new content is added as native, editable text.
Everything else -- title, five-step strip, rules, closing line, logo -- is
untouched.
"""
import shutil
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "deck.pptx"
OUT = "Pitch_Deck_AthenaRun_slide2-sources.pptx"
MEDIA = "ppt/media/image2.png"

PX = 6350  # EMU per pixel of the 1920x1080 export

BG      = (0x0E, 0x0E, 0x11)
WHITE   = RGBColor(0xEC, 0xEC, 0xEE)
MUTED   = RGBColor(0x9A, 0x9A, 0xA2)
ORANGE  = RGBColor(0xFF, 0x6B, 0x3D)
DIVIDER = (0x1F, 0x1F, 0x22)

SANS = "Arial"
MONO = "Courier New"

COLS = [110, 735, 1301]   # column left edges, measured off the export
DIVS = [677, 1243]        # vertical column dividers

STATS = [
    ("2500", "%", WHITE,  "more software defects by 2028"),
    ("25",   "%", ORANGE, "of planned AI spend deferred to 2027"),
    ("153",  "%", WHITE,  "more architectural design flaws"),
]
SOURCES = ("Sources: Gartner, Predicts 2026  ·  "
           "Forrester, Predictions 2026  ·  Apiiro, 2025")


# --- 1. erase the superseded regions from the slide image -------------------
with zipfile.ZipFile(SRC) as z:
    with z.open(MEDIA) as fh:
        img = Image.open(fh).convert("RGB")

d = ImageDraw.Draw(img)
d.rectangle([104, 545, 104 + 1711, 545 + 205], fill=BG)   # old stat band
d.rectangle([104, 820, 104 + 1000, 820 + 34], fill=BG)    # old source line
for x in DIVS:                                            # restore dividers
    d.rectangle([x, 571, x, 571 + 179], fill=DIVIDER)
img.save("image2_patched.png")

# --- 2. swap the image back into the package --------------------------------
with zipfile.ZipFile(SRC) as zin, zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = open("image2_patched.png", "rb").read() if item.filename == MEDIA \
            else zin.read(item.filename)
        zout.writestr(item, data)


# --- 3. add the new content as native text ----------------------------------
def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x * PX, y * PX, w * PX, h * PX)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return tf


def run(p, text, size, colour, font=SANS):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = colour
    r.font.name = font


prs = Presentation(OUT)
slide = list(prs.slides)[1]

for (num, pct, colour, label), left in zip(STATS, COLS):
    tf = textbox(slide, left, 560, 520, 110)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    run(tf.paragraphs[0], num, 50, colour)
    run(tf.paragraphs[0], pct, 20, MUTED)

    tf = textbox(slide, left, 673, 520, 40)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    run(tf.paragraphs[0], label, 11, MUTED)

tf = textbox(slide, 110, 824, 1300, 24)
tf.paragraphs[0].alignment = PP_ALIGN.LEFT
run(tf.paragraphs[0], SOURCES, 8, MUTED, font=MONO)

prs.save(OUT)
print(f"wrote {OUT}")
